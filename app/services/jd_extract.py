"""JD 文件文本抽取（同步；文件不落库）。

口径（PRD v1.5 §3.2.2）：
- 支持 `.pdf`（文本型）/ `.docx` / `.txt` / `.md`，单文件 ≤5MB
- **不做 OCR**：图片型 PDF / 扫描件明确报错，提示改用文本型 PDF 或手动粘贴
- 只抽纯文本，不做结构化拆分（结构化是下一步 C1 AI 拆解的事）
- **文件不落库、不进 COS**：抽取完即丢弃，只把纯文本回填到前端 JD 原文输入框
"""
from __future__ import annotations

import asyncio
import io
import logging
import re
from pathlib import PurePosixPath

from app.core.errors import ErrorCode, bad_request

logger = logging.getLogger(__name__)

MAX_BYTES = 5 * 1024 * 1024
ALLOWED_EXTS = {".pdf", ".docx", ".txt", ".md"}
# 扩展名可以伪造，用文件头再核一道
_MAGIC = {".pdf": b"%PDF", ".docx": b"PK\x03\x04"}

_UNSUPPORTED = "不支持的文件类型，仅支持文本型 PDF、DOCX、TXT、MD"
_NO_TEXT = "该文件无可提取文本，请改用文本型 PDF 或手动粘贴纯文本"


def _ext_of(filename: str, unsupported: str = ErrorCode.JD_FILE_UNSUPPORTED) -> str:
    ext = PurePosixPath(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise bad_request(unsupported, _UNSUPPORTED)
    return ext


def _clean(text: str) -> str:
    """压缩空白：去行尾空格、合并连续空行。"""
    lines = [re.sub(r"[ \t\u3000]+", " ", ln).strip() for ln in text.splitlines()]
    out: list[str] = []
    blank = 0
    for ln in lines:
        if not ln:
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        out.append(ln)
    return "\n".join(out).strip()


def _from_pdf(data: bytes) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    # JD 里的任职要求常写成表格，丢掉会漏掉大量内容
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


def _from_plain(data: bytes) -> str:
    for encoding in ("utf-8", "gbk", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


async def extract_text_async(
    filename: str,
    data: bytes,
    *,
    too_large: str = ErrorCode.JD_FILE_TOO_LARGE,
    unsupported: str = ErrorCode.JD_FILE_UNSUPPORTED,
    no_text: str = ErrorCode.JD_FILE_NO_TEXT,
) -> str:
    """在线程池里跑同步抽取：pdfplumber / python-docx 都是阻塞库，别堵事件循环。

    错误码可覆盖：简历（3.3）复用同一套抽取逻辑，但要报 `candidate.resume_*`，
    否则前端三语映射会指到职位那一类文案上。
    """
    return await asyncio.to_thread(
        extract_text,
        filename,
        data,
        too_large=too_large,
        unsupported=unsupported,
        no_text=no_text,
    )


def extract_text(
    filename: str,
    data: bytes,
    *,
    too_large: str = ErrorCode.JD_FILE_TOO_LARGE,
    unsupported: str = ErrorCode.JD_FILE_UNSUPPORTED,
    no_text: str = ErrorCode.JD_FILE_NO_TEXT,
) -> str:
    """同步抽取纯文本。失败（不支持 / 过大 / 无文本层）一律抛 400。"""
    ext = _ext_of(filename, unsupported)
    if len(data) > MAX_BYTES:
        raise bad_request(too_large, "文件超过 5MB，请精简后重试")
    # 纯文本没有文件头可核，跳过 magic 校验
    magic = _MAGIC.get(ext)
    if magic and not data.startswith(magic):
        raise bad_request(unsupported, f"文件内容与扩展名不符（{ext}）")

    try:
        if ext == ".pdf":
            text = _from_pdf(data)
        elif ext == ".docx":
            text = _from_docx(data)
        else:
            text = _from_plain(data)
    except Exception as exc:  # noqa: BLE001 —— 解析库异常统一转成可读错误
        logger.warning("文件解析失败（%s）：%s: %s", filename, type(exc).__name__, exc)
        raise bad_request(no_text, _NO_TEXT) from exc

    text = _clean(text)
    if not text:
        raise bad_request(no_text, _NO_TEXT)
    return text
